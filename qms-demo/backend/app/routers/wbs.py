from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session
from ..database import get_db
from .. import models, schemas
from ..auth import get_current_user, assert_project_access, allowed_section_ids, write_audit
from ..seed import SPECIALTY_TEMPLATES

router = APIRouter(prefix="/api/wbs", tags=["wbs"])


def _serialize(n: models.WbsNode) -> dict:
    return {
        "id": n.id,
        "section_id": n.section_id,
        "parent_id": n.parent_id,
        "name": n.name,
        "code": n.code,
        "node_type": n.node_type,
        "specialty": n.specialty,
        "path": n.path,
        "level": n.level,
        "sort_order": n.sort_order,
        "default_template_codes": n.default_template_codes or [],
    }


def _build_tree(nodes: list[models.WbsNode]) -> list[dict]:
    by_id = {n.id: {**_serialize(n), "children": []} for n in nodes}
    roots = []
    for n in nodes:
        item = by_id[n.id]
        if n.parent_id and n.parent_id in by_id:
            by_id[n.parent_id]["children"].append(item)
        else:
            roots.append(item)
    return roots


@router.get("/tree")
def wbs_tree(section_id: str, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    section = db.get(models.Section, section_id)
    if not section:
        raise HTTPException(404, "标段不存在")
    assert_project_access(db, user, section.project_id)
    allowed = allowed_section_ids(db, user, section.project_id)
    if allowed is not None and section_id not in allowed:
        raise HTTPException(403, "无该标段权限")
    nodes = (
        db.query(models.WbsNode)
        .filter(models.WbsNode.section_id == section_id)
        .order_by(models.WbsNode.level, models.WbsNode.sort_order, models.WbsNode.code)
        .all()
    )
    return _build_tree(nodes)


@router.post("")
def create_wbs(body: schemas.WbsCreate, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    section = db.get(models.Section, body.section_id)
    if not section:
        raise HTTPException(404, "标段不存在")
    assert_project_access(db, user, section.project_id, roles=["admin", "qc", "owner"])
    parent = db.get(models.WbsNode, body.parent_id) if body.parent_id else None
    level = (parent.level + 1) if parent else 1
    path = f"{parent.path}/{body.code}" if parent else body.code
    templates = body.default_template_codes or SPECIALTY_TEMPLATES.get(body.specialty, [])
    node = models.WbsNode(
        section_id=body.section_id,
        parent_id=body.parent_id,
        name=body.name,
        code=body.code,
        node_type=body.node_type,
        specialty=body.specialty,
        level=level,
        path=path,
        sort_order=body.sort_order,
        default_template_codes=templates,
    )
    db.add(node)
    write_audit(db, user, "wbs.create", f"{node.code} {node.name}")
    db.commit()
    db.refresh(node)
    return _serialize(node)


@router.patch("/{node_id}")
def update_wbs(node_id: str, body: schemas.WbsUpdate, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    node = db.get(models.WbsNode, node_id)
    if not node:
        raise HTTPException(404, "节点不存在")
    section = db.get(models.Section, node.section_id)
    assert_project_access(db, user, section.project_id, roles=["admin", "qc", "owner"])
    data = body.model_dump(exclude_unset=True)
    for k, v in data.items():
        setattr(node, k, v)
    if "parent_id" in data or "code" in data:
        parent = db.get(models.WbsNode, node.parent_id) if node.parent_id else None
        node.level = (parent.level + 1) if parent else 1
        node.path = f"{parent.path}/{node.code}" if parent else node.code
    write_audit(db, user, "wbs.update", node.code)
    db.commit()
    return _serialize(node)


@router.delete("/{node_id}")
def delete_wbs(node_id: str, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    node = db.get(models.WbsNode, node_id)
    if not node:
        raise HTTPException(404, "节点不存在")
    section = db.get(models.Section, node.section_id)
    assert_project_access(db, user, section.project_id, roles=["admin", "qc"])
    kids = db.query(models.WbsNode).filter(models.WbsNode.parent_id == node_id).count()
    if kids:
        raise HTTPException(400, "请先删除子节点")
    db.delete(node)
    write_audit(db, user, "wbs.delete", node.code)
    db.commit()
    return {"ok": True}


@router.get("/export")
def export_wbs(section_id: str, user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    section = db.get(models.Section, section_id)
    if not section:
        raise HTTPException(404, "标段不存在")
    assert_project_access(db, user, section.project_id)
    nodes = db.query(models.WbsNode).filter(models.WbsNode.section_id == section_id).order_by(models.WbsNode.path).all()
    wb = Workbook()
    ws = wb.active
    ws.title = "WBS"
    ws.append(["code", "name", "node_type", "specialty", "parent_code", "templates"])
    code_map = {n.id: n.code for n in nodes}
    for n in nodes:
        parent_code = code_map.get(n.parent_id, "") if n.parent_id else ""
        ws.append([n.code, n.name, n.node_type, n.specialty, parent_code, ",".join(n.default_template_codes or [])])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=wbs-{section.code}.xlsx"},
    )


@router.post("/import")
async def import_wbs(section_id: str, file: UploadFile = File(...), user: models.User = Depends(get_current_user), db: Session = Depends(get_db)):
    section = db.get(models.Section, section_id)
    if not section:
        raise HTTPException(404, "标段不存在")
    assert_project_access(db, user, section.project_id, roles=["admin", "qc"])
    content = await file.read()
    wb = load_workbook(BytesIO(content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "空文件")
    header = [str(c).strip() if c else "" for c in rows[0]]
    # expect code,name,node_type,specialty,parent_code,templates
    created = 0
    code_to_id = {n.code: n.id for n in db.query(models.WbsNode).filter(models.WbsNode.section_id == section_id).all()}
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        data = {header[i]: row[i] for i in range(min(len(header), len(row)))}
        code = str(data.get("code", "")).strip()
        name = str(data.get("name", "")).strip()
        if not code or not name:
            continue
        if code in code_to_id:
            continue
        parent_code = str(data.get("parent_code") or "").strip()
        parent_id = code_to_id.get(parent_code)
        parent = db.get(models.WbsNode, parent_id) if parent_id else None
        specialty = str(data.get("specialty") or "").strip()
        templates = [t for t in str(data.get("templates") or "").split(",") if t]
        if not templates:
            templates = SPECIALTY_TEMPLATES.get(specialty, [])
        node = models.WbsNode(
            section_id=section_id,
            parent_id=parent_id,
            name=name,
            code=code,
            node_type=str(data.get("node_type") or "item"),
            specialty=specialty,
            level=(parent.level + 1) if parent else 1,
            path=f"{parent.path}/{code}" if parent else code,
            default_template_codes=templates,
        )
        db.add(node)
        db.flush()
        code_to_id[code] = node.id
        created += 1
    write_audit(db, user, "wbs.import", f"section={section.code} created={created}")
    db.commit()
    return {"created": created}
